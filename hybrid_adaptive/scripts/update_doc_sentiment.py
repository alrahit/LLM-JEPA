import os
from docx import Document

# Path to your document
filepath = r"D:\LLM_JEPA\hybrid_adaptive\results\Research_Comparison_and_Results.docx"

if os.path.exists(filepath):
    doc = Document(filepath)
    print("Found existing document. Appending Rotten Tomatoes results...")
else:
    doc = Document()
    doc.add_heading('LinAlign: Experimental Results', 0)

# -------------------------
# Rotten Tomatoes Section
# -------------------------
doc.add_heading('7. Rotten Tomatoes (Sentiment Classification & Label Space Adherence)', level=1)
p = doc.add_paragraph()
p.add_run('This dataset uses non-standard sentiment labels ("Good" and "Bad"). The evaluation measures the model\'s ability to adhere strictly to this target label space rather than defaulting to instruction-tuned synonyms like "Positive" or "Negative".\n\n')

table = doc.add_table(rows=3, cols=3)
table.style = 'Light Shading Accent 1'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Dataset'
hdr_cells[1].text = 'Model / Configuration'
hdr_cells[2].text = 'True Label Accuracy'

data = [
    ('Rotten Tomatoes', 'Baseline (Raw Qwen2-1.5B)', '0.38% (4 / 1066)'),
    ('Rotten Tomatoes', 'LinAlign Model (Ours)', '91.84% (979 / 1066)')
]

for i, row in enumerate(data):
    cells = table.rows[i+1].cells
    cells[0].text = row[0]
    cells[1].text = row[1]
    cells[2].text = row[2]

obs = doc.add_paragraph()
obs.add_run('\nObservation (Label Space Adherence): ').bold = True
obs.add_run('The baseline model exhibits severe generative drift, utilizing the target labels ("Good"/"Bad") in only 1.1% (12/1066) of its outputs, defaulting instead to standard conversational synonyms. The LinAlign model achieves 91.84% accuracy, proving the dynamic InfoNCE controller successfully enforces strict adherence to arbitrary target label spaces without synonym hallucination.')

# Save the updated document
doc.save(filepath)
print(f"Rotten Tomatoes results successfully appended to: {filepath}")